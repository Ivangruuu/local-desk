import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from localdesk.documents import extract, excerpts, table_text


class Formats(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def read(self,name):
        return extract(self.root/name)

    def test_text_html_mail_rtf(self):
        (self.root/'text.txt').write_text('Проверка UTF-16: заказчик',encoding='utf-16')
        self.assertIn('заказчик',self.read('text.txt')['text'])
        (self.root/'web.html').write_text('<p>Visible</p><script>Hidden</script>')
        self.assertEqual(self.read('web.html')['text'],'Visible')
        (self.root/'mail.eml').write_text('Subject: Test\nContent-Type: text/plain; charset=utf-8\n\nHello mail')
        self.assertIn('Hello mail',self.read('mail.eml')['text'])
        (self.root/'text.rtf').write_text(r'{\rtf1\ansi Hello RTF}')
        self.assertIn('Hello RTF',self.read('text.rtf')['text'])

    def test_docx_pptx(self):
        from docx import Document
        doc=Document()
        doc.add_paragraph('Document alpha')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='Item'
        table.cell(0,1).text='42'
        doc.save(self.root/'doc.docx')
        text=self.read('doc.docx')['text']
        self.assertIn('Document alpha',text)
        self.assertIn('Item\t42',text)
        from pptx import Presentation
        ppt=Presentation()
        slide=ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text='Slide beta'
        ppt.save(self.root/'slides.pptx')
        self.assertIn('Slide beta',self.read('slides.pptx')['text'])

    def test_excel_csv_and_missing_formulas(self):
        import openpyxl
        book=openpyxl.Workbook()
        sheet=book.active
        sheet.append(['Компания','Выручка, ₽','Прибыль, ₽'])
        sheet.append(['А',100,20])
        sheet.append(['А',150,30])
        sheet.append(['Б',200,'=B4*0.2'])
        book.save(self.root/'table.xlsx')
        parsed=self.read('table.xlsx')
        sums=json.loads(parsed['summaries'][0])['company_totals']
        self.assertEqual(sums[0]['totals']['Выручка, ₽']['sum'],250)
        self.assertIsNone(sums[1]['totals']['Прибыль, ₽']['sum'])
        self.assertTrue(parsed['warnings'])
        (self.root/'table.csv').write_text('Код;Выручка\n00123;100,5\n00456;99,5\n')
        parsed=self.read('table.csv')
        self.assertIn('00123',parsed['text'])
        self.assertEqual(json.loads(parsed['summaries'][0])['numeric_statistics']['Выручка']['sum'],200)

    def test_pdf_and_ocr(self):
        import pymupdf
        pdf=pymupdf.open()
        page=pdf.new_page()
        page.insert_text((50,50),'PDF text gamma')
        pdf.save(self.root/'text.pdf')
        pdf.close()
        self.assertIn('PDF text gamma',self.read('text.pdf')['text'])
        from PIL import Image,ImageDraw,ImageFont
        image=Image.new('RGB',(1000,180),'white')
        ImageDraw.Draw(image).text((30,40),'LOCAL OCR 8427',fill='black',font=ImageFont.truetype('/System/Library/Fonts/Menlo.ttc',54))
        image.save(self.root/'scan.png')
        self.assertIn('8427',self.read('scan.png')['text'])
        pdf=pymupdf.open()
        page=pdf.new_page(width=1000,height=180)
        page.insert_image(page.rect,filename=str(self.root/'scan.png'))
        pdf.save(self.root/'scan.pdf')
        pdf.close()
        parsed=self.read('scan.pdf')
        self.assertIn('8427',parsed['text'])
        self.assertTrue(parsed['warnings'])

    def test_odt_epub_and_retrieval(self):
        from odf.opendocument import OpenDocumentText
        from odf.text import P
        doc=OpenDocumentText()
        doc.text.addElement(P(text='ODT delta'))
        doc.save(str(self.root/'file.odt'))
        self.assertIn('ODT delta',self.read('file.odt')['text'])
        with zipfile.ZipFile(self.root/'book.epub','w') as z:
            z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
            z.writestr('content.opf','<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="a" href="chapter.html"/></manifest><spine><itemref idref="a"/></spine></package>')
            z.writestr('chapter.html','<p>EPUB epsilon</p>')
        self.assertIn('EPUB epsilon',self.read('book.epub')['text'])
        doc={'text':'Filler text. '*2000+'Ошибка таймаута checkout-api. Код UNIQUE-8842.', 'summaries':[]}
        excerpt,coverage=excerpts(doc,'Ошибка таймаута checkout-api',4000)
        self.assertIn('UNIQUE-8842',excerpt)
        self.assertIn('Выдержки',coverage)


if __name__=='__main__':
    unittest.main(verbosity=2)
