"""Local Desk: a native Qt macOS chat. No web view or HTTP server."""
import html
import json
import os
import sys
import time
import uuid
from pathlib import Path

import markdown
from PySide6.QtCore import Qt, QProcess, QTimer, Signal
from PySide6.QtGui import QFont, QIcon, QTextDocument, QKeySequence, QShortcut
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QHBoxLayout, QVBoxLayout,
    QLabel, QPushButton, QListWidget, QListWidgetItem, QTextBrowser, QTextEdit,
    QComboBox, QCheckBox, QButtonGroup, QFileDialog, QFrame, QProgressBar,
    QMessageBox, QScrollArea)

from .models import ROOT, MODELS, ready
from .paths import ASSETS, process_command
from .worker import split_answer

HERE = Path(__file__).parent
STATE = ROOT / 'data' / 'chats.json'

STYLE = '''
QWidget { background: #161B26; color: #E9EDF5; font-family: ".AppleSystemUIFont"; font-size: 14px; }
QFrame#sidebar { background: #10141D; border-right: 1px solid #252D3D; }
QFrame#sidebar * { background: transparent; }
QLabel#brand { font-size: 24px; font-weight: 700; color: #FFFFFF; }
QLabel#muted { color: #8D9AB1; font-size: 12px; }
QLabel#title { font-size: 21px; font-weight: 600; }
QPushButton { background: #222B3B; border: 1px solid #344056; border-radius: 9px; padding: 9px 14px; }
QPushButton:hover { background: #2C384D; border-color: #6778A1; }
QPushButton:disabled { color: #58657C; border-color: #283043; background: #1C2330; }
QPushButton#primary { background: #7C8CFF; color: #0C1122; font-weight: 700; border: none; }
QPushButton#primary:hover { background: #99A5FF; }
QPushButton#new { background: #20283A; text-align: left; border-color: #3B4663; }
QPushButton#mode { background: #171D2A; border-color: #303B51; color: #93A1B7; padding: 8px 15px; }
QPushButton#mode:checked { background: #313B64; color: #DDE2FF; border-color: #6E7BD3; }
QListWidget { border: none; outline: none; background: #10141D; }
QListWidget::item { padding: 13px 10px; border-radius: 8px; margin: 3px 0; color: #AEBBD0; }
QListWidget::item:selected { background: #252F46; color: #F7F9FF; }
QListWidget::item:hover { background: #1B2231; }
QTextBrowser { border: none; background: #161B26; padding: 10px; selection-background-color: #5269AE; }
QFrame#composer { border: 1px solid #3C4860; border-radius: 15px; background: #1D2534; }
QFrame#composer QLabel, QFrame#composer QTextEdit { background: transparent; }
QTextEdit { border: none; padding: 6px; background: transparent; font-size: 15px; }
QComboBox { background: #222B3B; border: 1px solid #344056; border-radius: 8px; padding: 8px 14px; min-width: 185px; }
QComboBox::drop-down { border: none; width: 24px; }
QComboBox QAbstractItemView { background: #222B3B; selection-background-color: #3B4870; }
QCheckBox { color: #8D9AB1; spacing: 7px; font-size: 12px; }
QProgressBar { border: none; background: #202A3C; border-radius: 2px; max-height: 3px; }
QProgressBar::chunk { background: #8798FF; }
QScrollBar:vertical { background: transparent; width: 7px; }
QScrollBar::handle:vertical { background: #3B4861; border-radius: 3px; min-height: 30px; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical { background: transparent; }
'''


class Composer(QTextEdit):
    submit = Signal()
    files_dropped = Signal(list)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter) and not event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
            self.submit.emit()
        else:
            super().keyPressEvent(event)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragEnterEvent(event)

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragMoveEvent(event)

    def dropEvent(self, event):
        if event.mimeData().hasUrls():
            self.files_dropped.emit([u.toLocalFile() for u in event.mimeData().urls() if u.isLocalFile()])
            event.acceptProposedAction()
        else:
            super().dropEvent(event)


class Window(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle('Local Desk')
        self.resize(1180, 820)
        self.setMinimumSize(960, 650)
        self.setAcceptDrops(True)
        self.chats = self.read_state()
        self.current = None
        self.pending_files = []
        self.process = None
        self.buffer = b''
        self.busy = False
        self.family = 'qwen'
        self.thinking = False
        self.raw = ''
        self.active_model = None
        self.build_ui()
        self.render_timer = QTimer(self)
        self.render_timer.setInterval(80)
        self.render_timer.timeout.connect(self.update_stream)
        self.dirty = False
        self.render_timer.start()
        QShortcut(QKeySequence('Ctrl+N'), self, self.new_chat)
        QShortcut(QKeySequence('Ctrl+O'), self, self.pick_files)
        self.refresh_list()
        if self.chats:
            self.select_chat(self.chats[0]['id'])
        else:
            self.new_chat()

    def build_ui(self):
        root = QWidget()
        self.setCentralWidget(root)
        layout = QHBoxLayout(root)
        layout.setContentsMargins(0,0,0,0)
        layout.setSpacing(0)
        sidebar = QFrame()
        sidebar.setObjectName('sidebar')
        sidebar.setFixedWidth(225)
        side = QVBoxLayout(sidebar)
        side.setContentsMargins(18,28,18,22)
        brand = QLabel('Local Desk')
        brand.setObjectName('brand')
        side.addWidget(brand)
        subtitle = QLabel('Твой помощник на Mac')
        subtitle.setObjectName('muted')
        side.addWidget(subtitle)
        side.addSpacing(22)
        new = QPushButton('+  Новый чат')
        new.setObjectName('new')
        new.clicked.connect(self.new_chat)
        side.addWidget(new)
        side.addSpacing(15)
        self.chat_list = QListWidget()
        self.chat_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.chat_list.setTextElideMode(Qt.TextElideMode.ElideRight)
        self.chat_list.itemClicked.connect(lambda item: self.select_chat(item.data(Qt.ItemDataRole.UserRole)))
        side.addWidget(self.chat_list,1)
        delete = QPushButton('Удалить чат')
        delete.clicked.connect(self.delete_chat)
        side.addWidget(delete)
        unload = QPushButton('Освободить память')
        unload.clicked.connect(self.stop)
        side.addWidget(unload)
        side.addSpacing(12)
        local = QLabel('●  Всё остаётся на этом Mac')
        local.setObjectName('muted')
        side.addWidget(local)
        layout.addWidget(sidebar)
        main = QVBoxLayout()
        main.setContentsMargins(30,24,30,20)
        main.setSpacing(12)
        header = QHBoxLayout()
        title = QLabel('Давай разберёмся')
        title.setObjectName('title')
        header.addWidget(title,1)
        self.model_box = QComboBox()
        for key, spec in MODELS.items():
            self.model_box.addItem(spec['name'],key)
            self.model_box.setItemData(self.model_box.count()-1,spec['description'],Qt.ItemDataRole.ToolTipRole)
        self.model_box.currentIndexChanged.connect(self.model_changed)
        header.addWidget(self.model_box)
        self.models_button = QPushButton('Модели')
        self.models_button.clicked.connect(self.manage_models)
        header.addWidget(self.models_button)
        main.addLayout(header)
        self.model_detail = QLabel()
        self.model_detail.setObjectName('muted')
        main.addWidget(self.model_detail)
        modes = QHBoxLayout()
        self.fast = QPushButton('Быстро')
        self.think = QPushButton('С рассуждением')
        group = QButtonGroup(self)
        for button in (self.fast,self.think):
            button.setObjectName('mode')
            button.setCheckable(True)
            group.addButton(button)
            modes.addWidget(button)
        self.fast.setChecked(True)
        self.think.setToolTip('Модель тратит больше времени на проверку и рассуждение.')
        modes.addStretch()
        self.show_thought = QCheckBox('Показать рассуждение')
        self.show_thought.toggled.connect(self.render_chat)
        modes.addWidget(self.show_thought)
        main.addLayout(modes)
        self.browser = QTextBrowser()
        self.browser.setOpenExternalLinks(False)
        self.browser.setAcceptDrops(False)
        self.browser.setFont(QFont('.AppleSystemUIFont',15))
        main.addWidget(self.browser,1)
        self.file_row = QHBoxLayout()
        self.file_row.setSpacing(5)
        main.addLayout(self.file_row)
        composer = QFrame()
        composer.setObjectName('composer')
        compose_layout = QVBoxLayout(composer)
        compose_layout.setContentsMargins(12,8,12,12)
        self.input = Composer()
        self.input.setFixedHeight(85)
        self.input.setPlaceholderText('Напиши сообщение или перетащи сюда файлы…')
        self.input.submit.connect(self.send)
        self.input.files_dropped.connect(self.add_files)
        compose_layout.addWidget(self.input)
        actions = QHBoxLayout()
        self.attach_button = QPushButton('+  Файлы')
        self.attach_button.clicked.connect(self.pick_files)
        actions.addWidget(self.attach_button)
        hint = QLabel('Enter — отправить · Shift+Enter — новая строка')
        hint.setObjectName('muted')
        actions.addWidget(hint,1)
        self.send_button = QPushButton('Отправить ↑')
        self.send_button.setObjectName('primary')
        self.send_button.clicked.connect(self.send_or_stop)
        actions.addWidget(self.send_button)
        compose_layout.addLayout(actions)
        main.addWidget(composer)
        self.progress = QProgressBar()
        self.progress.setRange(0,0)
        self.progress.hide()
        main.addWidget(self.progress)
        footer = QHBoxLayout()
        self.status = QLabel('Готов к работе')
        self.status.setObjectName('muted')
        self.status.setWordWrap(True)
        footer.addWidget(self.status,1)
        copy = QPushButton('Копировать ответ')
        copy.clicked.connect(self.copy_answer)
        footer.addWidget(copy)
        save = QPushButton('Сохранить ответ')
        save.clicked.connect(self.save_answer)
        footer.addWidget(save)
        main.addLayout(footer)
        layout.addLayout(main,1)
        self.model_changed()

    def read_state(self):
        try:
            data = json.loads(STATE.read_text(encoding='utf-8'))
            if isinstance(data,list) and all(isinstance(c,dict) and 'id' in c and 'messages' in c for c in data):
                return data
        except (OSError,ValueError):
            pass
        return []

    def persist(self):
        STATE.parent.mkdir(parents=True, exist_ok=True)
        temp = STATE.with_suffix('.tmp')
        temp.write_text(json.dumps(self.chats,ensure_ascii=False,indent=2),encoding='utf-8')
        temp.replace(STATE)

    def chat(self):
        return next((c for c in self.chats if c['id']==self.current),None)

    def refresh_list(self):
        self.chat_list.clear()
        for chat in self.chats:
            item=QListWidgetItem(chat['title'])
            item.setData(Qt.ItemDataRole.UserRole,chat['id'])
            self.chat_list.addItem(item)
            if chat['id']==self.current:
                self.chat_list.setCurrentItem(item)

    def new_chat(self):
        if self.busy:
            self.stop()
        chat={'id':uuid.uuid4().hex,'title':'Новый чат','messages':[],'model':self.model_box.currentData()}
        self.chats.insert(0,chat)
        self.select_chat(chat['id'])
        self.persist()

    def select_chat(self, chat_id):
        if self.busy:
            self.stop()
        self.current=chat_id
        self.pending_files=[]
        chat=self.chat()
        if chat:
            index=self.model_box.findData(chat.get('model','qwen9'))
            self.model_box.setCurrentIndex(max(0,index))
        self.refresh_list()
        self.refresh_files()
        self.render_chat()
        self.input.setFocus()

    def delete_chat(self):
        if not self.chat():
            return
        choice=QMessageBox.question(self,'Удалить чат?','Локальная история этого чата будет удалена.')
        if choice!=QMessageBox.StandardButton.Yes:
            return
        self.stop()
        self.chats=[c for c in self.chats if c['id']!=self.current]
        if self.chats:
            self.select_chat(self.chats[0]['id'])
        else:
            self.new_chat()
        self.persist()

    def model_changed(self):
        key=self.model_box.currentData()
        if not key:
            return
        if self.active_model and key!=self.active_model:
            self.release_worker()
        spec=MODELS[key]
        self.model_detail.setText(spec['description'] + (' · установлена' if ready(key) else ' · не установлена — открой «Модели»'))
        if self.current and self.chat():
            self.chat()['model']=key
            self.persist()

    def manage_models(self):
        from .model_manager import ModelManager
        self.release_worker()
        dialog = ModelManager(self.model_box.currentData(), self)
        dialog.exec()
        self.model_changed()

    def pick_files(self):
        files,_=QFileDialog.getOpenFileNames(self,'Выбери документы',str(Path.home()/'Documents'),
            'Все файлы (*);;Документы (*.pdf *.docx *.doc *.odt *.rtf *.txt *.md);;Таблицы (*.xlsx *.xls *.xlsb *.csv *.tsv *.ods);;Презентации (*.pptx *.odp)')
        self.add_files(files)

    def add_files(self,files):
        if self.busy:
            return
        for path in files:
            if Path(path).is_file() and path not in self.pending_files:
                self.pending_files.append(path)
        self.refresh_files()

    def refresh_files(self):
        while self.file_row.count():
            item=self.file_row.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        for path in self.pending_files:
            chip=QPushButton(Path(path).name+'  ×')
            chip.setToolTip(path)
            chip.clicked.connect(lambda checked=False,p=path:self.remove_file(p))
            self.file_row.addWidget(chip)
        self.file_row.addStretch()

    def remove_file(self,path):
        self.pending_files.remove(path)
        self.refresh_files()

    def dragEnterEvent(self,event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self,event):
        self.add_files([u.toLocalFile() for u in event.mimeData().urls() if u.isLocalFile()])
        event.acceptProposedAction()

    def set_busy(self,busy):
        self.busy=busy
        self.send_button.setText('Остановить ■' if busy else 'Отправить ↑')
        self.model_box.setEnabled(not busy)
        self.models_button.setEnabled(not busy)
        self.fast.setEnabled(not busy)
        self.think.setEnabled(not busy)
        self.attach_button.setEnabled(not busy)
        self.progress.setVisible(busy)

    def send_or_stop(self):
        self.stop() if self.busy else self.send()

    def send(self):
        if self.busy:
            return
        text=self.input.toPlainText().strip()
        if not text and not self.pending_files:
            return
        key=self.model_box.currentData()
        if not ready(key):
            self.manage_models()
            return
        text=text or 'Объясни содержимое прикреплённых файлов.'
        chat=self.chat()
        files=list(self.pending_files)
        chat['messages'].append({'role':'user','content':text,'files':files})
        if chat['title']=='Новый чат':
            chat['title']=text.replace('\n',' ')[:26]
        model_messages=[{'role':m['role'],'content':m['content']} for m in chat['messages'] if m.get('content') and not m.get('failed')]
        all_files=list(dict.fromkeys(p for m in chat['messages'] for p in m.get('files',[])))
        self.family=MODELS[key]['family']
        self.thinking=self.think.isChecked()
        self.raw=''
        chat['messages'].append({'role':'assistant','content':'','reasoning':'','model':MODELS[key]['name']})
        self.pending_files=[]
        self.input.clear()
        self.refresh_files()
        self.refresh_list()
        self.render_chat()
        self.set_busy(True)
        self.status.setText('Подготавливаю модель…')
        self.persist()
        if self.process is None or self.process.state()==QProcess.ProcessState.NotRunning:
            self.start_worker(key)
        request={'model':key,'messages':model_messages,'paths':all_files,'thinking':self.thinking}
        self.process.write((json.dumps(request,ensure_ascii=False)+'\n').encode())

    def start_worker(self,key):
        process=QProcess(self)
        process.setWorkingDirectory(str(ROOT))
        process.setProcessChannelMode(QProcess.ProcessChannelMode.SeparateChannels)
        process.readyReadStandardOutput.connect(self.read_worker)
        process.readyReadStandardError.connect(self.read_diagnostics)
        process.finished.connect(self.worker_finished)
        self.buffer=b''
        self.process=process
        self.active_model=key
        program, arguments = process_command('--worker')
        process.start(program, arguments)
        if not process.waitForStarted(3000):
            self.fail('Не удалось запустить процесс модели.')

    def read_diagnostics(self):
        if self.process:
            text=bytes(self.process.readAllStandardError())
            if text:
                log=ROOT/'logs'/'worker.log'
                log.parent.mkdir(parents=True, exist_ok=True)
                with log.open('ab') as handle:
                    handle.write(text)

    def read_worker(self):
        if not self.process:
            return
        self.buffer+=bytes(self.process.readAllStandardOutput())
        while b'\n' in self.buffer:
            line,self.buffer=self.buffer.split(b'\n',1)
            try:
                event=json.loads(line)
            except ValueError:
                continue
            self.handle_event(event)

    def handle_event(self,event):
        kind=event['event']
        if kind=='status':
            self.status.setText(event['text'])
        elif kind=='documents' and self.chat():
            message=self.chat()['messages'][-2]
            message['files']=event['paths']
            message['document_notes']=event['notes']
            if event['trimmed']:
                message['context_note']='Самые старые сообщения не поместились в контекст модели.'
            self.render_chat()
        elif kind=='token':
            self.raw+=event['text']
            self.dirty=True
        elif kind=='done' and self.chat():
            message=self.chat()['messages'][-1]
            message['content']=event['answer']
            message['reasoning']=event['reasoning']
            message['metrics']=f'{event["tps"]:.1f} токен/с · {event["seconds"]:.1f} с · MLX {event["memory"]:.1f} ГБ'
            if event.get('finish')=='length':
                message['metrics']+=' · лимит ответа достигнут'
            self.dirty=False
            self.set_busy(False)
            self.status.setText(message['metrics'])
            self.persist()
            self.render_chat()
        elif kind=='error':
            self.fail(event['text'])

    def update_stream(self):
        if self.dirty and self.busy and self.chat():
            answer,reasoning=split_answer(self.raw,self.thinking,self.family)
            self.chat()['messages'][-1].update(content=answer,reasoning=reasoning)
            self.render_chat()
            self.dirty=False

    def fail(self,text):
        if self.chat() and self.chat()['messages'] and self.chat()['messages'][-1]['role']=='assistant':
            self.chat()['messages'][-1].update(content='Не удалось выполнить запрос: '+text,failed=True)
        self.set_busy(False)
        self.status.setText(text)
        self.persist()
        self.render_chat()

    def worker_finished(self,code,status):
        if self.busy:
            self.fail('Процесс модели завершился. Можно повторить запрос или выбрать более лёгкую модель.')
        self.active_model=None

    def release_worker(self):
        if self.process:
            process=self.process
            self.process=None
            process.blockSignals(True)
            process.kill()
            process.waitForFinished(3000)
            process.deleteLater()
        self.active_model=None

    def stop(self):
        if self.busy and self.chat():
            self.update_stream()
            message=self.chat()['messages'][-1]
            if not message.get('content'):
                message['failed']=True
                message['content']='Генерация остановлена.'
            message['metrics']='Остановлено пользователем'
        self.set_busy(False)
        self.release_worker()
        self.status.setText('Остановлено · память модели освобождена')
        self.persist()
        self.render_chat()

    def render_chat(self,*args):
        if not self.chat():
            return
        bar=self.browser.verticalScrollBar()
        at_bottom=bar.maximum()-bar.value()<40
        old_value=bar.value()
        messages=self.chat()['messages']
        if not messages:
            self.browser.setHtml('''<div style="margin:70px 15px"><p style="font-size:32px;color:#EEF2FA">С чем помочь?</p>
                <p style="font-size:16px;color:#97A5BC">Разобрать документ, написать код или сформулировать мысль.</p>
                <p style="font-size:14px;color:#73839E">Прикрепи файл кнопкой «Файлы» или перетащи его в окно.<br>Можно указать путь прямо в сообщении.</p></div>''')
            return
        blocks=[]
        for message in messages:
            user=message['role']=='user'
            label='ВЫ' if user else html.escape(message.get('model','ПОМОЩНИК'))
            color='#AEBBFF' if user else '#79DCC4'
            content=message.get('content','')
            if user:
                body=html.escape(content).replace('\n','<br>')
            else:
                # Disable raw HTML from documents/model output in the text renderer.
                renderer=markdown.Markdown(extensions=['fenced_code','tables','sane_lists'])
                renderer.preprocessors.deregister('html_block')
                renderer.inlinePatterns.deregister('html')
                body=renderer.convert(content)
                if not content:
                    body='<span style="color:#8D9AB1">Рассуждаю…</span>' if message.get('reasoning') else '<span style="color:#8D9AB1">Готовлю ответ…</span>'
            thought=''
            if self.show_thought.isChecked() and message.get('reasoning'):
                thought='<p style="color:#8090A9;font-size:12px">РАССУЖДЕНИЕ</p><p style="color:#8D9AB1;font-size:13px">'+html.escape(message['reasoning']).replace('\n','<br>')+'</p>'
            notes=[]
            for doc in message.get('document_notes',[]):
                notes.append(doc['name']+' · '+doc['coverage'])
                notes.extend(doc.get('warnings',[]))
            if not message.get('document_notes'):
                notes.extend(Path(p).name for p in message.get('files',[]))
            if message.get('context_note'):
                notes.append(message['context_note'])
            if message.get('metrics'):
                notes.append(message['metrics'])
            meta='<br>'.join(html.escape(n) for n in notes)
            blocks.append(f'<p style="color:{color};font-size:11px;font-weight:bold;margin-top:26px">{label}</p>{thought}<div style="font-size:15px;line-height:1.5">{body}</div><p style="color:#7D8CA4;font-size:11px">{meta}</p><hr color="#293346">')
        self.browser.document().setDefaultStyleSheet('pre {background:#101621;white-space:pre-wrap;padding:14px;color:#C7D7EF;} code {font-family:Menlo;font-size:13px;} table {border-collapse:collapse;} td,th {padding:6px;border:1px solid #34405A;} a {color:#AEBBFF;}')
        self.browser.setHtml('<html><body style="color:#E0E6F0">'+''.join(blocks)+'</body></html>')
        bar.setValue(bar.maximum() if at_bottom else old_value)

    def last_answer(self):
        if self.chat():
            return next((m.get('content','') for m in reversed(self.chat()['messages']) if m['role']=='assistant' and not m.get('failed')), '')
        return ''

    def copy_answer(self):
        QApplication.clipboard().setText(self.last_answer())
        self.status.setText('Ответ скопирован')

    def save_answer(self):
        text=self.last_answer()
        if not text:
            self.status.setText('Сначала получи ответ модели.')
            return
        filename,_=QFileDialog.getSaveFileName(self,'Сохранить ответ',str(Path.home()/'Documents'/'Ответ.txt'),
            'Текст (*.txt);;Markdown (*.md);;JSON (*.json);;Word (*.docx);;PDF (*.pdf)')
        if not filename:
            return
        try:
            path=Path(filename)
            if path.suffix.lower()=='.json':
                raw=text.strip()
                if raw.startswith('```'):
                    raw='\n'.join(raw.splitlines()[1:-1])
                path.write_text(json.dumps(json.loads(raw),ensure_ascii=False,indent=2),encoding='utf-8')
            elif path.suffix.lower()=='.docx':
                from docx import Document
                doc=Document()
                for paragraph in text.split('\n\n'):
                    doc.add_paragraph(paragraph)
                doc.save(path)
            elif path.suffix.lower()=='.pdf':
                from PySide6.QtPrintSupport import QPrinter
                printer=QPrinter(QPrinter.PrinterMode.HighResolution)
                printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
                printer.setOutputFileName(str(path))
                document=QTextDocument()
                document.setPlainText(text)
                document.print_(printer)
            else:
                path.write_text(text,encoding='utf-8')
            self.status.setText('Сохранено: '+path.name)
        except Exception as error:
            QMessageBox.warning(self,'Не удалось сохранить',str(error))

    def closeEvent(self,event):
        self.stop()
        event.accept()


def main():
    app=QApplication(sys.argv)
    app.setApplicationName('Local Desk')
    app.setOrganizationName('LocalDesk')
    app.setStyle('Fusion')
    app.setStyleSheet(STYLE)
    icon=ASSETS/'icon.png'
    if icon.exists():
        app.setWindowIcon(QIcon(str(icon)))
    window=Window()
    app.aboutToQuit.connect(window.stop)
    window.show()
    if '--smoke-test' in sys.argv:
        QTimer.singleShot(700,lambda: window.grab().save(str(ROOT/'ui-preview.png')))
        QTimer.singleShot(1000,app.quit)
    sys.exit(app.exec())


if __name__=='__main__':
    main()
