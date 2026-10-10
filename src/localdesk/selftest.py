"""No-user-data diagnostics suitable for a relocated release bundle."""
import json
import tempfile
from pathlib import Path


def main():
    import mlx.core as mx
    cpu_only = '--cpu-only' in __import__('sys').argv
    if cpu_only:
        mx.set_default_device(mx.cpu)
    from .documents import extract, ocr
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw, ImageFont
    import pymupdf
    from .paths import ASSETS
    from .worker import split_answer
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        note = root / 'note.txt'
        note.write_text('Release check ALPHA-8127', encoding='utf-8')
        assert 'ALPHA-8127' in extract(note)['text']
        document = Document()
        document.add_paragraph('Document ALPHA-8127')
        document.save(root / 'note.docx')
        assert 'ALPHA-8127' in extract(root / 'note.docx')['text']
        book = Workbook()
        book.active.append(['Company', 'revenue'])
        book.active.append(['Synthetic', 123])
        book.save(root / 'note.xlsx')
        assert '123' in extract(root / 'note.xlsx')['text']
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((40, 80), 'PDF ALPHA-8127')
        pdf.save(root / 'note.pdf')
        pdf.close()
        assert 'ALPHA-8127' in extract(root / 'note.pdf')['text']
        image = Image.new('RGB', (1200, 220), 'white')
        draw = ImageDraw.Draw(image)
        draw.text((40, 60), 'CHECK ALPHA 8127', fill='black', font=ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc', 58))
        image.save(root / 'ocr.png')
        assert '8127' in ocr(root / 'ocr.png')
        mx.eval(mx.arange(4) + 1)
        assert mx.sum(mx.arange(4)).item() == 6
        assert split_answer('hidden</think>ok', True, 'qwen')[0] == 'ok'
        if '--network-check' in __import__('sys').argv:
            from huggingface_hub import hf_hub_download
            from .models import MODELS
            spec = MODELS['qwen9']
            config = hf_hub_download(spec['repo'], 'config.json', revision=spec['revision'], local_dir=root / 'hub')
            assert json.loads(Path(config).read_text())['model_type'] == 'qwen3_5'
    print(json.dumps({'ok': True, 'frozen': __import__('sys').frozen if hasattr(__import__('sys'), 'frozen') else False,
                      'checks': ['mlx-cpu' if cpu_only else 'mlx-metal', 'txt', 'docx', 'xlsx', 'pdf', 'native-ocr', 'reasoning'],
                      'ocr_bundled': (ASSETS / 'ocr').is_file()}), flush=True)
